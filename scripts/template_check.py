"""检查 templates/<模板名>/ 能否给 Deck 用，并列出各版式页的形状位置，供填写 template.json。

用法：
    python template_check.py <模板名>

检查项：template.json 字段齐全；template.pptx 正好 4 页（封面、目录、过渡页、内容页）；
Deck 要填写的形状都在、名字对；没有残留 p:custDataLst 形状标签（复制后会让 PowerPoint 拒绝打开）；
正文区（body）在页面内、在内容页标题之下。有问题时退出码为 1。
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent))
from hust_deck import CM, Template  # noqa: E402

PAGES = ["封面", "目录", "过渡页", "内容页"]
REQUIRED = [{"封面标题", "汇报人", "日期"}, {"目录条目"}, {"章节标题"}, {"页面标题", "章节标签"}]


def walk(shapes, depth=0):
    for sh in shapes:
        yield depth, sh
        if sh.shape_type == 6:
            yield from walk(sh.shapes, depth + 1)


def cm(v):
    return f"{v / CM:6.2f}" if v is not None else "     -"


def main():
    t = Template(sys.argv[1])
    prs = Presentation(str(t.pptx))
    sw, sh_ = prs.slide_width / CM, prs.slide_height / CM
    problems = []
    slides = list(prs.slides)
    if len(slides) != 4:
        problems.append(f"template.pptx 有 {len(slides)} 页，应为 4 页：{'、'.join(PAGES)}")
    print(f"模板「{t.name}」（{t.title}），页面 {sw:.2f} × {sh_:.2f} cm\n")
    for i, (slide, page, need) in enumerate(zip(slides, PAGES, REQUIRED), 1):
        print(f"== 第 {i} 页：{page}    名称 / 左 / 上 / 宽 / 高（cm）/ 文字")
        names = set()
        for depth, s in walk(slide.shapes):
            names.add(s.name)
            text = s.text_frame.text.replace("\n", " ")[:16] if s.has_text_frame else ""
            mark = "  ← Deck 填写" if s.name in need else ""
            print(f"  {'  ' * depth}{s.name:<16}{cm(s.left)}{cm(s.top)}{cm(s.width)}{cm(s.height)}  {text}{mark}")
        for name in sorted(need - names):
            problems.append(f"{page}缺少名为「{name}」的形状（在 PowerPoint 的选择窗格里改名）")
        if any(True for _ in slide._element.iter(qn("p:custDataLst"))):
            problems.append(f"{page}有 p:custDataLst 形状标签，删掉（见 references/templates.md）")
        print()

    body_t, body_b = t.y0 / CM, t.y0 / CM + t.h
    body_l, body_r = t.x0 / CM, t.x0 / CM + t.w
    print(f"正文区 body：x {body_l:.2f}–{body_r:.2f}、y {body_t:.2f}–{body_b:.2f} cm（从页面左上角量）")
    if body_l < 0 or body_r > sw + 0.01 or body_t < 0 or body_b > sh_ + 0.01:
        problems.append("正文区超出页面，改 template.json 的 body")
    if len(slides) == 4:
        title = next((s for _, s in walk(slides[3].shapes) if s.name == "页面标题"), None)
        if title is not None and (title.top + title.height) / CM > body_t + 0.01:
            problems.append(f"正文区上边 {body_t:.2f}cm 高于页面标题下沿 {(title.top + title.height) / CM:.2f}cm")

    print("\n" + ("\n".join(f"❌ {p}" for p in problems) if problems else "✅ 可以使用"))
    sys.exit(1 if problems else 0)


if __name__ == "__main__":
    main()
