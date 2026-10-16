"""学术汇报 PPT 的程序化质检。

用法：
    python qa_check.py <deck.pptx> [--source 材料1.md 材料2.tex ...] [--template 模板名]

模板默认取 Deck.save 记在文件属性（类别）里的模板名；正文区位置、主色、占位词、模板自带的装饰形状都从这套模板读。

检查项（❌ 必须修，⚠️ 逐条确认）：
    ❌ 同一页形状 ID 重复、多个形状共用同一个标签部件（PowerPoint 会拒绝打开整个文件）
    ❌ 形状超出页面
    ❌ 模板占位词残留（汇报标题、章节名、页面结论句……）
    ❌ 字号小于 10.5pt
    ⚠️ 文字估算宽度或行数超出文本框（按 tw 估算，中文按 1em；表格单元格不估算，以渲染图为准）
    ⚠️ 正文里有小于 13pt 的文字、正文字数超过 150
    ⚠️ 两个带文字的形状互相重叠
    ⚠️ 正文没有铺满正文区：内容上沿离正文区顶端超过 3cm，或下沿离底端超过 3.5cm（只看外框，卡片内部的空白以渲染图为准）
    ⚠️ 内容页没有口播稿备注
    ⚠️ 正文图片分辨率低于 150dpi（放映时发虚）、网图（web_figure）没有图注
    ⚠️ 页面标题过短，可能是主题标签而不是结论句；页面标题、章节标签会折成两行
    ⚠️ 标题、正文、口播稿里有 AI 套话（值得注意的是、综上所述、不是……而是……等）
    ⚠️ 近似错位：两个主要块（卡片、底块、图片、表格、图表）的同一条边相差 0.02–0.25cm，要么对齐要么拉开
    ⚠️ 同一构图（主要块的位置和尺寸）出现在 5 页及以上，版式节奏单调
    ⚠️ 页面上的数字在 --source 材料中找不到
    ⚠️ 颜色不在 hust_deck 配色里、含义色超过 3 种（有图例时另加图例里的类型数）、字号超过 4 档、
       同一标签重复 3 次以上
最后按顺序列出全部页面标题，供「只读标题」测试。有 ❌ 时退出码为 1。
"""
import argparse
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

# 版式页上由 Deck 填写的形状；它们在模板里的原文就是占位词，生成后还留着就是漏填
FILLED = {"封面标题", "汇报人", "日期", "目录条目", "章节标题", "页面标题", "章节标签"}
# 版式页上的固定文本框；「文本框 7 / 20」是 9.22 原始 PPT 里的旧名称，直接改旧 deck 时也能识别
TITLE_NAMES = {"页面标题", "文本框 7"}
HEADER = TITLE_NAMES | {"章节标签", "文本框 20", "封面标题", "汇报人", "日期", "章节标题"}
CM = 360000


sys.path.insert(0, str(Path(__file__).resolve().parent))
from hust_deck import GRAY, INK, LINE, TEXT, TONES, TYPE_NAMES, WHITE, templates, tw, use_template  # noqa: E402

TYPE_LABELS = tuple(TYPE_NAMES.values()) + ("定律或定理", "公式或方程", "应用", "方法")

# 每页最多出现的「含义色」种数（不含中性色）：主色、强调色、错误色加上图例里的类型色
MAX_TONES = 3
# 正文内容离正文区上、下边超过这些距离（cm）时提示大块空白
GAP_TOP, GAP_BOTTOM = 3.0, 3.5
# 正文图片按放置尺寸计算的最低分辨率
MIN_DPI = 150
# 主要块：参与近似错位和构图重复检查的形状（按 hust_deck 的命名前缀）
BLOCKS = ("卡片", "底色", "图片", "网图", "表格", "图表", "截图")
NEAR_MISS = (0.02, 0.25)  # cm
# 构图签名只看面积不小于这个值（cm²）的块；同一签名出现这么多页时提示
SIG_MIN_AREA, SIG_MAX_REPEAT = 20, 5
# AI 套话：句式类，不含「系统」「全面」这类在技术文本里常有正当用法的词
SLOP = ["众所周知", "不可否认", "毋庸置疑", "毫无疑问", "值得注意的是", "需要指出的是", "更重要的是", "说白了",
        "归根结底", "从某种意义上", "底层逻辑", "一句话总结", "一句话讲透", "综上所述", "总而言之", "总的来说",
        "具有重要意义", "深远影响", "拭目以待", "赋能", "抓手", "降维打击"]
SLOP_RE = re.compile("|".join(map(re.escape, SLOP)) + r"|不是[^，。；！？\n]{1,15}而是")


def tones_of(colors):
    """颜色属于哪些 tone（中性色不计）。"""
    return {name for name, t in TONES.items() if name != "neutral" and colors & {t.base, t.tint, t.ink}}


def rgb(color):
    """显式 RGB 颜色的 hex；主题色（模板自带的形状常用）返回 None。"""
    return str(color.rgb) if color.type == MSO_COLOR_TYPE.RGB else None


def frames(sh):
    """形状里的全部文本框：普通形状一个，表格每个单元格一个。"""
    if sh.has_text_frame:
        return [sh.text_frame]
    if getattr(sh, "has_table", False) and sh.has_table:
        return [cell.text_frame for row in sh.table.rows for cell in row.cells]
    return []


def colors_of(sh):
    """形状用到的颜色：填充、边框、文字。"""
    out = set()
    fill = getattr(sh, "fill", None)
    if fill is not None and fill.type == MSO_FILL.SOLID:
        out.add(rgb(fill.fore_color))
    line = getattr(sh, "line", None)
    if line is not None and line.fill.type == MSO_FILL.SOLID:
        out.add(rgb(line.color))
    for tf in frames(sh):
        for p in tf.paragraphs:
            for r in p.runs:
                if r.text.strip() and r.font.color.type is not None:
                    out.add(rgb(r.font.color))
    out.discard(None)
    return out


def iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from iter_shapes(sh.shapes)
        else:
            yield sh


def text_of(sh):
    return "\n".join(tf.text for tf in frames(sh))


def sizes_of(sh):
    out = []
    for tf in frames(sh):
        for p in tf.paragraphs:
            for r in p.runs:
                if r.font.size and r.text.strip():
                    out.append((r.font.size.pt, r.text))
    return out


def overflow(sh):
    """估算文字是否超出文本框，返回问题描述或 None。"""
    tf = sh.text_frame
    w = (sh.width - tf.margin_left - tf.margin_right) / CM
    h = (sh.height - tf.margin_top - tf.margin_bottom) / CM
    if w <= 0:
        return None
    total_h = 0.0
    for p in tf.paragraphs:
        runs = [r for r in p.runs if r.text]
        if not runs:
            continue
        size = max((r.font.size.pt if r.font.size else 18) for r in runs)
        width = sum(tw(r.text, r.font.size.pt if r.font.size else 18) for r in runs)
        indent = (int(p._p.pPr.get("marL", 0)) / CM) if p._p.pPr is not None else 0
        if not tf.word_wrap:
            if width > w + 0.15:
                return f"单行宽约 {width:.1f}cm，框宽 {w:.1f}cm"
            lines = 1
        else:
            lines = max(1, -(-width // max(w - indent, 0.1)))
        lh = p.line_spacing.pt if p.line_spacing is not None and hasattr(p.line_spacing, "pt") else size * 1.25
        total_h += lines * lh * 0.03528 + (p.space_before.pt * 0.03528 if p.space_before else 0)
    if tf.word_wrap and total_h > h + 0.25:
        return f"文字高约 {total_h:.1f}cm，框高 {h:.1f}cm"
    return None


def bbox(sh):
    return (sh.left, sh.top, sh.left + sh.width, sh.top + sh.height)


def overlap(a, b):
    x = min(a[2], b[2]) - max(a[0], b[0])
    y = min(a[3], b[3]) - max(a[1], b[1])
    return x > CM * 0.05 and y > CM * 0.05


def contains(a, b):
    return a[0] <= b[0] and a[1] <= b[1] and a[2] >= b[2] and a[3] >= b[3]


def load_template(prs, name):
    """载入 deck 所用的模板，返回 (模板, 占位词, 模板自带的装饰形状名)。"""
    if name is None:
        m = re.fullmatch(r"template:(\S+)", prs.core_properties.category or "")
        if not m:
            sys.exit(f"文件属性里没有记录模板（不是 Deck 生成的？），用 --template 指定：{'、'.join(templates())}")
        name = m.group(1)
    t = use_template(name)
    stale, decor = set(), set()
    for slide in Presentation(str(t.pptx)).slides:
        for sh in iter_shapes(slide.shapes):
            if sh.name in FILLED:
                stale.add(sh.text_frame.text.strip())
            else:  # 顶栏、校徽、标题线等，不计入正文样式统计，也允许伸出页外
                decor.add(sh.name)
    return t, stale, decor


def source_numbers(paths):
    text = "\n".join(Path(p).read_text(encoding="utf-8") for p in paths)
    return set(re.findall(r"\d+(?:\.\d+)?", text))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--source", nargs="*", default=[], help="核对数字用的原始材料（md / tex / txt / json）")
    ap.add_argument("--template", help="模板名；默认取 deck 文件属性里记录的模板")
    args = ap.parse_args()

    prs = Presentation(args.deck)
    tmpl, stale_words, decor = load_template(prs, args.template)
    # 规定的配色；正文出现其他颜色就是手写了颜色，应改用 tone
    palette = {INK, TEXT, GRAY, LINE, WHITE, "000000"} | {c for t in TONES.values() for c in (t.base, t.tint, t.ink)}
    W, H = prs.slide_width, prs.slide_height
    nums = source_numbers(args.source) if args.source else None
    errors = warns = 0
    titles = []
    layouts = {}

    def report(level, i, msg):
        nonlocal errors, warns
        if level == "❌":
            errors += 1
        else:
            warns += 1
        print(f"{level} 第 {i} 页：{msg}")

    for i, slide in enumerate(prs.slides, 1):
        ids = [e.get("id") for e in slide._element.iter(qn("p:cNvPr"))]
        dup = sorted({x for x in ids if ids.count(x) > 1}, key=int)
        if dup:  # LibreOffice 和格式校验都放过，但 PowerPoint 会拒绝打开整个文件
            report("❌", i, f"形状 ID 重复：{', '.join(dup)}")
        tags = [e.get(qn("r:id")) for e in slide._element.iter(qn("p:tags"))]
        shared = sorted({x for x in tags if tags.count(x) > 1})
        if shared:  # 复制带标签的形状时会出现；同样会让 PowerPoint 打不开文件
            report("❌", i, f"多个形状共用同一个标签部件：{', '.join(shared)}")
        shapes = list(iter_shapes(slide.shapes))
        names = {sh.name for sh in shapes}
        is_content = bool(names & TITLE_NAMES) and i > 1
        body_chars = 0
        body_sizes = []
        texted = []
        colors = set()
        chips = {}
        n_legend = 0
        body_boxes = []
        blocks = []
        texts = []
        for sh in shapes:
            if sh.left is None:
                continue
            in_body = is_content and sh.name not in HEADER | decor and not sh.name.startswith("灯片编号")
            if in_body:
                body_boxes.append(bbox(sh))
                if sh.name.startswith(BLOCKS):
                    blocks.append(sh)
                if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    dpi = sh.image.size[0] / (sh.width / 914400)
                    if dpi < MIN_DPI:
                        report("⚠️", i, f"「{sh.name}」按放置尺寸只有约 {dpi:.0f}dpi，放映时发虚，换大图或缩小")
                colors |= colors_of(sh)
                # 图例，以及写着类型名的标签、标题条、表格单元格，都算是对类型色的说明
                if sh.name.startswith("图例"):
                    n_legend += 1
                elif sh.name.startswith(("标签", "标题条", "表格")):
                    n_legend += sum(tf.text.lstrip("●· ").startswith(TYPE_LABELS) for tf in frames(sh))
                if sh.name.startswith("标签"):
                    chips[text_of(sh).strip()] = chips.get(text_of(sh).strip(), 0) + 1
            l, t, r, b = bbox(sh)
            # 模板自带的形状（如华科的校徽图片）本来就可能有一部分在页外
            if sh.name not in HEADER | decor and not sh.name.startswith("灯片编号") and (
                    l < -CM * 0.05 or t < -CM * 0.05 or r > W + CM * 0.05 or b > H + CM * 0.05):
                report("❌", i, f"「{sh.name}」超出页面")
            txt = text_of(sh).strip()
            if not txt:
                continue
            for w in stale_words:
                if w in txt:
                    report("❌", i, f"「{sh.name}」残留模板占位词「{w}」")
            for size, run in sizes_of(sh):
                if size < 10.5:
                    report("❌", i, f"「{sh.name}」字号 {size:g}pt 太小：{run[:20]}")
            msg = overflow(sh) if sh.has_text_frame else None
            if msg:
                report("⚠️", i, f"「{sh.name}」可能溢出（{msg}）：{txt[:24]}")
            if is_content and sh.name not in HEADER and not sh.name.startswith("灯片编号"):
                body_chars += len(re.sub(r"\s", "", txt))
                body_sizes += [s for s, _ in sizes_of(sh)]
                texts.append(txt)
                if sh.has_text_frame:
                    texted.append(sh)
            if nums is not None and i > 1 and not sh.name.startswith("灯片编号"):
                for n in re.findall(r"\d+(?:\.\d+)?", txt):
                    if len(n.replace(".", "")) >= 2 and not re.fullmatch(r"0\d", n) and n not in nums:
                        report("⚠️", i, f"数字 {n} 在材料中找不到（「{sh.name}」：{txt[:24]}）")
        for a_i, a in enumerate(texted):
            for b in texted[a_i + 1:]:
                ba, bb = bbox(a), bbox(b)
                if overlap(ba, bb) and not contains(ba, bb) and not contains(bb, ba):
                    report("⚠️", i, f"「{a.name}」与「{b.name}」重叠")
        if is_content:
            title = _title(shapes)
            titles.append((i, title))
            if len(title) < 10:
                report("⚠️", i, f"标题「{title}」偏短，检查是否写成了结论句")
            for sh in shapes:  # 模板里这两个框会自动换行，估算宽度超出就是折成了两行
                if sh.name in TITLE_NAMES | {"章节标签", "文本框 20"}:
                    run = sh.text_frame.paragraphs[0].runs[0]
                    width = tw(sh.text_frame.text.strip(), run.font.size.pt if run.font.size else 24)
                    if width > sh.width / CM - 0.5:
                        report("⚠️", i, f"「{sh.name}」约 {width:.1f}cm，框宽 {sh.width / CM - 0.5:.1f}cm，会折成两行")
            small = sorted({s for s in body_sizes if s < 13})
            if small:
                report("⚠️", i, f"有小于 13pt 的文字（{small}），正文 16、标签和图注 13")
            stray = sorted(colors - palette)
            if stray:
                report("⚠️", i, f"颜色不在 hust_deck 配色里（改用 tone）：{' '.join(stray)}")
            used = tones_of(colors)
            if len(used) > MAX_TONES + n_legend:
                report("⚠️", i, f"用了 {len(used)} 种含义色（{'、'.join(sorted(used))}），类型色要配图例")
            if len(set(body_sizes)) > 4:
                report("⚠️", i, f"正文字号 {len(set(body_sizes))} 档（上限 4）：{sorted(set(body_sizes))}")
            for label, k in chips.items():
                if k >= 3:
                    report("⚠️", i, f"标签「{label}」重复 {k} 次，改成分组标题")
            n_web = sum(n.startswith("网图") for n in names)
            n_cap = sum(n.startswith("图注") for n in names)
            if n_web > n_cap:
                report("⚠️", i, f"{n_web} 张网图只有 {n_cap} 条图注，每张图下面写清图里是什么、标号的含义")
            if body_boxes:
                top = (min(b[1] for b in body_boxes) - tmpl.y0) / CM
                bottom = tmpl.h - (max(b[3] for b in body_boxes) - tmpl.y0) / CM
                if top > GAP_TOP or bottom > GAP_BOTTOM:
                    report("⚠️", i, f"正文区上方空 {top:.1f}cm、下方空 {bottom:.1f}cm，内容没有铺满，调整布局")
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
            slop = sorted({m.group(0) for t in [title, notes, *texts] for m in SLOP_RE.finditer(t)})
            if slop:
                report("⚠️", i, f"有 AI 套话：{'、'.join(slop)}，删掉或换成具体说法")
            misses = near_misses(blocks)
            if misses:
                report("⚠️", i, f"近似错位（要么对齐要么拉开）：{'；'.join(misses[:3])}")
            sig = tuple(sorted((sh.name.split()[0], *(round(v / CM) for v in (sh.left, sh.top, sh.width, sh.height)))
                               for sh in blocks if sh.width * sh.height / CM ** 2 >= SIG_MIN_AREA))
            if sig:
                layouts.setdefault(sig, []).append(i)
            if body_chars > 150:
                report("⚠️", i, f"正文约 {body_chars} 字，超过 150 字，考虑删减或拆页")
            if not notes:
                report("⚠️", i, "没有口播稿备注")

    for pages in layouts.values():
        if len(pages) >= SIG_MAX_REPEAT:
            report("⚠️", pages[0], f"第 {'、'.join(map(str, pages))} 页构图相同，换几页的排法（图为主、全宽流程、表格等）")

    print("\n只读标题测试（按顺序读，应能连成完整的论证）：")
    for i, t in titles:
        print(f"  {i:>2}. {t}")
    print(f"\n❌ {errors} 项，⚠️ {warns} 项")
    sys.exit(1 if errors else 0)


def near_misses(blocks):
    """主要块之间同一条边（左、右、上、下）相差 NEAR_MISS 范围内的情况。"""
    out = []
    edges = ("左", lambda b: b[0]), ("上", lambda b: b[1]), ("右", lambda b: b[2]), ("下", lambda b: b[3])
    for k, a in enumerate(blocks):
        for b in blocks[k + 1:]:
            for side, f in edges:
                d = abs(f(bbox(a)) - f(bbox(b))) / CM
                if NEAR_MISS[0] < d < NEAR_MISS[1]:
                    out.append(f"「{a.name}」和「{b.name}」{side}边差 {d:.2f}cm")
    return out


def _title(shapes):
    for sh in shapes:
        if sh.name in TITLE_NAMES:
            return sh.text_frame.text.strip()
    return ""


if __name__ == "__main__":
    main()
