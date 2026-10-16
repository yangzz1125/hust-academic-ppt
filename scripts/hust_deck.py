"""学术汇报 PPT 的生成库：在 templates/<模板>/template.pptx 上克隆版式页，正文用原生形状绘制。

用法（在每份汇报自己的生成脚本里）：

    import sys
    sys.path.insert(0, r"<skill>/scripts")   # <skill> 换成本 skill 目录的绝对路径
    from hust_deck import *

    deck = Deck("hust")                      # 模板名 = templates/ 下的目录名
    deck.cover("汇报标题", "汇报人", "2026年10月7日")
    deck.toc(["研究背景", "方法设计", "后续计划"])
    c = deck.content("研究背景", "现有课程问答缺少可追溯的知识依据", notes="口播稿……")
    c.card(Grid().box(0, 0, 6, 6))
    ...
    deck.save("out.pptx")

页面按调用顺序排列；模板自带的 4 页在保存时删除。

正文画布：内容页标题线下方的正文区域，原点在左下角、y 轴向上，单位 cm。
位置和宽高由 template.json 的 body 给出（华科模板为宽 31cm、高 14.3cm），所有 Canvas 方法都用这个坐标。
一个进程里同时只用一套模板：创建 Deck 时载入，画布换算、Grid 和 TONES["primary"] 都随之改变。

配色按含义使用，方法只接受「色调」tone，不接受任意颜色：
    "primary"  模板主色（华科蓝 0259A0）：页面主体、卡片标题条、序号、主要箭头
    "accent"   橙色：本次重点、产出、关键结论
    "neutral"  灰色：已完成、次要、无方向
    "error"    红色：只表示错误
    类型色     "concept" "law" "quantity" "formula" "application" "method"：
               六类知识点各一种颜色，全篇一致，出现时配图例或在表格里写明类型
每个 tone 有三种用法：base（实色）、tint（浅底）、ink（浅底上的文字色）。
"""
import copy
import json
import struct
import zlib
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_DATA_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Pt

__all__ = ["Deck", "Canvas", "Box", "Grid", "P", "R", "var", "place", "tw", "pt2cm", "fill_text",
           "INK", "TEXT", "GRAY", "LINE", "WHITE", "TONES", "TYPE_NAMES", "S_TITLE", "S_BODY", "S_SMALL",
           "MSO_SHAPE", "templates"]

TEMPLATES = Path(__file__).resolve().parent.parent / "templates"

# ---- 配色（正文只用这些）。中性色固定；有含义的颜色在 TONES 里，primary 由模板给出
INK, TEXT, GRAY, LINE, WHITE = "1E293B", "334155", "64748B", "D7E0EC", "FFFFFF"


class Tone:
    """一种有含义的颜色：base 实色，tint 浅底（base 与白色 12% 混合），ink 浅底上的文字色。"""

    def __init__(self, base, ink=None, tint=None):
        self.base = base
        self.ink = ink or base
        self.tint = tint or "".join(f"{round(c * 0.12 + 255 * 0.88):02X}" for c in bytes.fromhex(base))


# 模板没有改写时的色调；template.json 的 tones 必须给出 primary，也可以改写这里的任意一项
BASE_TONES = {
    "accent": Tone("E8710A", ink="C2410C"),
    "neutral": Tone(GRAY, tint="F1F5F9"),
    "error": Tone("C00000"),
    # 六类知识点
    "concept": Tone("2563EB"),
    "law": Tone("7C3AED"),
    "quantity": Tone("0F8B8D"),
    "formula": Tone("DB2777"),
    "application": Tone("15803D"),
    "method": Tone("475569"),
}
# 当前模板的全部色调。创建 Deck 时原地更新，所以 from hust_deck import * 拿到的 TONES 也会跟着变
TONES = {}
TYPE_NAMES = {"concept": "概念", "law": "定律", "quantity": "物理量", "formula": "公式",
              "application": "工程应用", "method": "分析方法"}

# ---- 字号（pt）：页面标题 24 由模板提供；正文只用这三档
S_TITLE, S_BODY, S_SMALL = 18, 16, 13

FONT = "微软雅黑"
LATIN_MATH = "Times New Roman"

CM = 360000


class Template:
    """templates/<name>/ 下的一套模板：template.pptx（封面、目录、过渡页、内容页 4 页）和 template.json。
    json 各字段的含义见 references/templates.md。"""

    def __init__(self, name):
        d = TEMPLATES / name
        if not (d / "template.json").exists():
            raise ValueError(f"没有模板「{name}」，可选：{'、'.join(templates())}")
        spec = json.loads((d / "template.json").read_text(encoding="utf-8"))
        self.name, self.pptx = name, d / "template.pptx"
        self.title = spec["title"]
        self.tones = spec["tones"]
        body = spec["body"]
        self.x0, self.y0 = round(body["x"] * CM), round(body["y"] * CM)  # 正文区左上角（EMU）
        self.w, self.h = body["w"], body["h"]  # 正文区宽高（cm）
        self.toc = spec["toc"]
        self.cover_title = spec["cover_title"]


def templates():
    """可用的模板名（templates/ 下带 template.json 的目录）。"""
    return sorted(p.parent.name for p in TEMPLATES.glob("*/template.json"))


_current = None


def use_template(name):
    """载入模板并设为当前模板：画布换算、Grid、TONES 都改用它。Deck 和 qa_check 调用。"""
    global _current
    t = Template(name)
    TONES.clear()
    TONES.update(BASE_TONES)
    for k, v in t.tones.items():
        TONES[k] = Tone(v) if isinstance(v, str) else Tone(**v)
    _current = t
    return t


def current():
    if _current is None:
        raise RuntimeError("还没有选模板：先创建 Deck(模板名)")
    return _current


def cn_num(k):
    """1–99 的中文数字：1→一，10→十，12→十二，20→二十，35→三十五。"""
    digits = "零一二三四五六七八九"
    tens, ones = divmod(k, 10)
    return (digits[tens] if tens > 1 else "") + ("十" if tens else "") + (digits[ones] if ones else "")


def ex(x):
    return Emu(round(current().x0 + x * CM))


def ey(y):
    t = current()
    return Emu(round(t.y0 + (t.h - y) * CM))


def el(d):
    return Emu(round(d * CM))


def pt2cm(v):
    return v * 0.03528


def tw(s, size):
    """估算单行文字宽度（cm）：汉字和全角符号按 1em，空格 0.3em，其他字符 0.55em。"""
    em = sum(1.0 if ord(ch) > 0x2E80 else 0.3 if ch == " " else 0.55 for ch in s)
    return pt2cm(em * size)


class Box:
    """画布上的矩形：l 为左边，t 为上边（y 轴向上），w、h 为宽高。"""

    def __init__(self, l, t, w, h):
        self.l, self.t, self.w, self.h = l, t, w, h

    cx = property(lambda s: s.l + s.w / 2)
    cy = property(lambda s: s.t - s.h / 2)
    r = property(lambda s: s.l + s.w)
    b = property(lambda s: s.t - s.h)
    west = property(lambda s: (s.l, s.cy))
    east = property(lambda s: (s.r, s.cy))
    north = property(lambda s: (s.cx, s.t))
    south = property(lambda s: (s.cx, s.b))


class Grid:
    """正文区的排版网格：cols 列 × rows 行，格间距统一为 gap（cm）。
    块的四条边都落在网格线上，所以不同大小的卡片随意组合（便当布局）也能自然对齐、间距一致。
    左右各留 0.6，上边留 0.5；下边 y 1.6，bottombar=True 时留到 y 2.2 给总结条。
    华科模板（31 × 14.3）上即 x 0.6–30.4、y 1.6–13.8。要在 Deck 创建之后用。

        g = Grid()
        c.card(g.box(0, 0, 8, 6))      # 第 0 列、第 0 行起，占 8 列 6 行：图为主
        c.card(g.box(8, 0, 4, 6))      # 右侧窄栏
    """

    def __init__(self, cols=12, rows=6, gap=0.6, bottombar=False):
        t = current()
        self.cols, self.rows, self.gap = cols, rows, gap
        self.l, self.t = 0.6, t.h - 0.5
        bottom = 2.2 if bottombar else 1.6
        self.cw = (t.w - 0.6 - self.l - gap * (cols - 1)) / cols
        self.rh = (self.t - bottom - gap * (rows - 1)) / rows

    def box(self, col, row, w=1, h=1):
        """第 col 列、第 row 行（都从 0 起，行从上往下数）开始，占 w 列 h 行的块。"""
        if col + w > self.cols or row + h > self.rows:
            raise ValueError(f"块 ({col}, {row}, {w}, {h}) 超出 {self.cols}×{self.rows} 网格")
        return Box(self.l + col * (self.cw + self.gap), self.t - row * (self.rh + self.gap),
                   w * self.cw + (w - 1) * self.gap, h * self.rh + (h - 1) * self.gap)


def place(anchor, x, y, w, h):
    """按 TikZ 风格的锚点（c / n / s / w / e / nw / ne / sw / se）把 w×h 的矩形放到 (x, y)。"""
    dx = {"w": 0, "e": -w}.get(anchor[-1], -w / 2)
    dy = {"n": 0, "s": h}.get(anchor[0], h / 2)
    return Box(x + dx, y + dy, w, h)


class R:
    """一段文字及其格式；未给出的格式沿用段落或文本框的默认值。
    sub / sup 为下标、上标；latin 指定西文字体（公式变量用 Times New Roman 加 italic）。"""

    def __init__(self, text, color=None, bold=None, size=None, italic=False, latin=None, sub=False, sup=False):
        self.text, self.color, self.bold, self.size = text, color, bold, size
        self.italic, self.latin, self.sub, self.sup = italic, latin, sub, sup


class P:
    """一个段落：若干 R 或字符串。bullet 为项目符号颜色，before 为段前距（pt）。"""

    def __init__(self, *runs, bullet=None, before=0):
        self.runs = [r if isinstance(r, R) else R(r) for r in runs]
        self.bullet, self.before = bullet, before


def var(sym, sub=None):
    """公式变量：Times New Roman 斜体，可带正体下标。返回 R 列表，放进 P(*var(...))。"""
    runs = [R(sym, italic=True, latin=LATIN_MATH)]
    if sub:
        runs.append(R(sub, latin=LATIN_MATH, sub=True))
    return runs


SHADOW = (f'<a:effectLst {nsdecls("a")}><a:outerShdw blurRad="50800" dist="19050" dir="5400000" '
          'algn="t" rotWithShape="0"><a:srgbClr val="000000"><a:alpha val="10000"/></a:srgbClr>'
          '</a:outerShdw></a:effectLst>')
ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
VANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def style_run(r, size, color, bold, italic=False, latin=None, sub=False, sup=False):
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    if italic:
        f.italic = True
    f.color.rgb = RGBColor.from_string(color)
    rPr = r._r.get_or_add_rPr()
    rPr.set("lang", "zh-CN")  # 不写时 PowerPoint 不按中文规则断行，行首会出现「，」
    rPr.set("altLang", "en-US")
    if sub or sup:
        rPr.set("baseline", "-25000" if sub else "30000")
    for tag in ("a:latin", "a:ea"):
        old = rPr.find(qn(tag))
        if old is not None:
            rPr.remove(old)
    etree.SubElement(rPr, qn("a:latin")).set("typeface", latin or FONT)
    etree.SubElement(rPr, qn("a:ea")).set("typeface", FONT)


def fill_text(tf, paras, size=S_BODY, color=TEXT, bold=False, align="l", line=None, anchor="t", wrap=True,
              margin=(0, 0, 0, 0)):
    """把 paras（str / P / R 或它们的列表）写进文本框，清掉原有内容。"""
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = VANCHOR[anchor]
    tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = (el(m) for m in margin)
    if isinstance(paras, (str, P, R)):
        paras = [paras]
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)
    for i, para in enumerate(paras):
        para = para if isinstance(para, P) else P(para)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        if line:
            p.line_spacing = Pt(line)
        if para.before:
            p.space_before = Pt(para.before)
        if para.bullet:
            pPr = p._p.get_or_add_pPr()
            ind = el(pt2cm(size) * 1.25)
            pPr.set("marL", str(ind))
            pPr.set("indent", str(-ind))
            pPr.append(parse_xml(f'<a:buClr {nsdecls("a")}><a:srgbClr val="{para.bullet}"/></a:buClr>'))
            pPr.append(parse_xml(f'<a:buSzPct {nsdecls("a")} val="80000"/>'))
            pPr.append(parse_xml(f'<a:buFont {nsdecls("a")} typeface="{FONT}"/>'))
            pPr.append(parse_xml(f'<a:buChar {nsdecls("a")} char="●"/>'))
        for rs in para.runs:
            r = p.add_run()
            r.text = rs.text
            style_run(r, rs.size or size, rs.color or color, bold if rs.bold is None else rs.bold,
                      rs.italic, rs.latin, rs.sub, rs.sup)


def plain(paras):
    """paras 的纯文本，每段一项。"""
    plist = [paras] if isinstance(paras, (str, P, R)) else paras
    out = []
    for p in plist:
        if isinstance(p, str):
            out.append(p)
        elif isinstance(p, R):
            out.append(p.text)
        else:
            out.append("".join(r.text for r in p.runs))
    return out


# ============================================================ 正文绘制

class Canvas:
    """一个内容页的正文区域。每个方法画一个原生对象并按「类型 序号」命名。"""

    def __init__(self, slide):
        self.slide = slide
        self.shapes = slide.shapes
        self.n = 0

    def _finish(self, shp, kind):
        # 自选形状自带主题样式（p:style），会带阴影和主题色；删掉后颜色全部由本模块显式写出
        style = shp._element.find(qn("p:style"))
        if style is not None:
            shp._element.remove(style)
        self.n += 1
        shp.name = f"{kind} {self.n}"
        return shp

    # ---- 基本图形
    def shape(self, kind, prst, box, fill=None, line=None, lw=0.8, radius=None, shadow=False):
        """底层方法：fill、line 只应取本模块的配色常量。"""
        shp = self._finish(self.shapes.add_shape(prst, ex(box.l), ey(box.t), el(box.w), el(box.h)), kind)
        if radius is not None:
            shp.adjustments[0] = min(0.5, radius / min(box.w, box.h))
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = RGBColor.from_string(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = RGBColor.from_string(line)
            shp.line.width = Pt(lw)
        else:
            shp.line.fill.background()
        if shadow:
            shp._element.spPr.append(parse_xml(SHADOW))
        return shp

    def card(self, box, header=None, tone="primary", solid=True, highlight=False, header_h=1.1):
        """白底圆角卡片，带浅阴影。
        header 给出时在顶部画标题条：solid=True 为实色白字，False 为浅底深色字。
        highlight=True 时边框改为 tone 的实色粗线，标出本页重点。
        返回标题条以下的正文区域 Box（没有标题条时返回卡片本身）。"""
        t = TONES[tone]
        self.shape("卡片", MSO_SHAPE.ROUNDED_RECTANGLE, box, WHITE,
                   t.base if highlight else LINE, 1.75 if highlight else 0.8, radius=0.247, shadow=True)
        if header is None:
            return box
        bar = Box(box.l, box.t, box.w, header_h)
        shp = self.shape("标题条", MSO_SHAPE.ROUND_2_SAME_RECTANGLE, bar, t.base if solid else t.tint,
                         radius=min(0.247, header_h / 2))
        shp.adjustments[1] = 0
        fill_text(shp.text_frame, header, size=S_TITLE, color=WHITE if solid else t.ink, bold=True,
                  anchor="m", wrap=False, margin=(0.5, 0.3, 0, 0))
        return Box(box.l, box.t - header_h, box.w, box.h - header_h)

    def tint(self, box, tone="primary"):
        """浅色底块。"""
        return self.shape("底色", MSO_SHAPE.ROUNDED_RECTANGLE, box, TONES[tone].tint, radius=0.247)

    # ---- 文字
    def text(self, box, paras, kind="文字", fill=None, **kw):
        """在 box 里放文本框；kw 传给 fill_text（size、color、bold、align、line、anchor、wrap）。"""
        tb = self._finish(self.shapes.add_textbox(ex(box.l), ey(box.t), el(box.w), el(box.h)), kind)
        if fill:
            tb.fill.solid()
            tb.fill.fore_color.rgb = RGBColor.from_string(fill)
        fill_text(tb.text_frame, paras, **kw)
        return tb

    def at(self, anchor, x, y, paras, size=S_BODY, w=None, h=None, line=None, **kw):
        """按锚点放文字。不给 w 时按估算宽度放单行、不换行；给 w 时在该宽度内换行。"""
        lines = plain(paras)
        wrap = w is not None
        w = w if wrap else max(tw(s, size) for s in lines) + 0.3
        h = h or pt2cm(line or size * 1.3) * len(lines)
        horiz = {"w": "l", "e": "r"}.get(anchor[-1], "c")
        vert = {"n": "t", "s": "b"}.get(anchor[0], "m")
        return self.text(place(anchor, x, y, w, h), paras, size=size, line=line, wrap=wrap,
                         align=kw.pop("align", horiz), anchor=vert, **kw)

    def title(self, anchor, x, y, s, **kw):
        """卡片或分组的标题：18pt 粗体深灰。"""
        return self.at(anchor, x, y, s, size=S_TITLE, color=INK, bold=True, **kw)

    def bullets(self, box, items, tone="primary", size=S_BODY, line=None, anchor="t", gap=0):
        """项目符号列表：items 为字符串或 P；圆点颜色取 tone。
        卡片较高、条目较少时用 anchor="m" 纵向居中、gap（pt）拉开条目间距，不让要点挤在上半部。"""
        dot = TONES[tone].base
        paras = [P(*(it.runs if isinstance(it, P) else [it]), bullet=dot, before=gap if k else 0)
                 for k, it in enumerate(items)]
        return self.text(box, paras, size=size, line=line or size * 1.5, anchor=anchor)

    def num(self, x, y, label, tone="primary", d=0.9):
        """实心圆序号。只在内容有先后顺序时使用。"""
        shp = self.shape("序号", MSO_SHAPE.OVAL, place("c", x, y, d, d), TONES[tone].base)
        fill_text(shp.text_frame, label, size=S_SMALL, color=WHITE, bold=True, align="c", anchor="m", wrap=False)
        return shp

    def chip(self, anchor, x, y, label, tone="primary", solid=False):
        """小标签（13pt 粗体）：默认浅底深色字，solid=True 为实色白字。返回标签的 Box。
        同一页不要重复同一个标签，改成分组标题。"""
        t = TONES[tone]
        fill, color = (t.base, WHITE) if solid else (t.tint, t.ink)
        box = place(anchor, x, y, tw(label, S_SMALL) + 0.5, pt2cm(S_SMALL * 1.05 + 8))
        shp = self.shape("标签", MSO_SHAPE.ROUNDED_RECTANGLE, box, fill, radius=0.141)
        fill_text(shp.text_frame, label, size=S_SMALL, color=color, bold=True, align="c", anchor="m", wrap=False)
        return box

    def node(self, x, y, paras, tone="primary", w=None, h=1.15, size=S_BODY):
        """图中的节点：tone 的浅底 + 实色边框。知识点用类型色（如 tone="law"）。返回 Box。"""
        first = plain(paras)[0]
        box = place("c", x, y, w or max(3.4, tw(first, size) + 0.6), h)
        t = TONES[tone]
        shp = self.shape("节点", MSO_SHAPE.ROUNDED_RECTANGLE, box, t.tint, t.base, 1.25, radius=0.176)
        fill_text(shp.text_frame, paras, size=size, color=TEXT, align="c", anchor="m", wrap=False)
        return box

    def heading(self, card, label, n=None, dy=0.9, tone="primary"):
        """卡片左上角的标题；给 n 时前面加序号圆点。"""
        y = card.t - dy
        if n is None:
            self.title("w", card.l + 0.6, y, label)
        else:
            self.num(card.l + 0.9, y, n, tone)
            self.title("w", card.l + 1.6, y, label)

    def keynum(self, x, y, value, label, tone="accent"):
        """关键数字：18pt 粗体的数值（默认强调色），正下方 13pt 灰色说明，整组以 (x, y) 为中心。
        用于一页里要听众记住的两三个数字，如结论页的「数字带」、效率对比。"""
        self.at("s", x, y + 0.4, value, size=S_TITLE, color=TONES[tone].ink, bold=True, kind="关键数字")
        self.at("n", x, y + 0.3, label, size=S_SMALL, color=GRAY, kind="关键数字")

    def output(self, card, text, h=1.8, tone="accent"):
        """卡片底部的「产出」框：默认橙色浅底，13pt 标签 + 13pt 粗体内容。"""
        box = Box(card.l + 0.3, card.b + 0.3 + h, card.w - 0.6, h)
        self.tint(box, tone)
        self.text(Box(box.l + 0.3, box.t - 0.2, box.w - 0.5, box.h - 0.35),
                  [R("产出", color=TONES[tone].ink, bold=True), P(R(text, color=INK, bold=True), before=2)],
                  size=S_SMALL, anchor="m")
        return box

    def bottombar(self, label, text, y=0.35):
        """页底总结条：左侧粗体标签，分隔线，右侧一句话。左右边和 Grid 的卡片对齐（左右各留 0.6）。"""
        bar = Box(0.6, y + 1.25, current().w - 1.2, 1.25)
        self.tint(bar)
        self.at("w", bar.l + 0.5, bar.cy, label, color=TONES["primary"].base, bold=True)
        sep = bar.l + 0.5 + tw(label, S_BODY) + 0.6
        self.line((sep, bar.cy - 0.35), (sep, bar.cy + 0.35), LINE, 1.0)
        self.at("w", sep + 0.4, bar.cy, text)

    def legend(self, x, y, tones, directed=True, undirected=True):
        """图例：tones 为类型色名列表（如 ["concept", "law"]），名称取 TYPE_NAMES；再加有方向 / 无方向两种线。"""
        for tone in tones:
            name = TYPE_NAMES[tone]
            t = TONES[tone]
            self.shape("图例", MSO_SHAPE.ROUNDED_RECTANGLE, place("c", x, y, 0.7, 0.42), t.tint, t.base, 1.25,
                       radius=0.07)
            self.at("w", x + 0.5, y, name, size=S_SMALL)
            x += 1.2 + tw(name, S_SMALL)
        if directed:
            self.line((x, y), (x + 1.2, y), GRAY, 1.1, arrow=True)
            self.at("w", x + 1.4, y, "有方向", size=S_SMALL)
            x += 3.4
        if undirected:
            self.line((x, y), (x + 1.2, y), GRAY, 1.1, dash=True)
            self.at("w", x + 1.4, y, "无方向", size=S_SMALL)

    # ---- 线条
    def _ln(self, shp, color, lw, arrow, dash, cap):
        shp.line.color.rgb = RGBColor.from_string(color)
        shp.line.width = Pt(lw)
        if dash:
            shp.line.dash_style = MSO_LINE.DASH
        ln = shp._element.spPr.find(qn("a:ln"))
        if cap:
            ln.set("cap", "rnd")
        if arrow:
            ln.append(parse_xml(f'<a:tailEnd {nsdecls("a")} type="stealth" w="lg" len="lg"/>'))

    def line(self, p1, p2, color=GRAY, lw=1.1, arrow=False, dash=False, cap=False):
        shp = self._finish(self.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, ex(p1[0]), ey(p1[1]),
                                                     ex(p2[0]), ey(p2[1])), "连线")
        self._ln(shp, color, lw, arrow, dash, cap)
        return shp

    def poly(self, pts, color=GRAY, lw=1.1, arrow=True, dash=False):
        """折线（如流程图里绕行的箭头），pts 为画布坐标点列。"""
        fb = self.shapes.build_freeform(ex(pts[0][0]), ey(pts[0][1]), scale=1.0)
        fb.add_line_segments([(ex(x), ey(y)) for x, y in pts[1:]], close=False)
        shp = self._finish(fb.convert_to_shape(), "折线")
        shp.fill.background()
        self._ln(shp, color, lw, arrow, dash, False)
        return shp

    def label(self, x, y, s):
        """连线上的说明文字（13pt），白底盖住线条。"""
        box = place("c", x, y, tw(s, S_SMALL) + 0.25, pt2cm(S_SMALL * 1.35))
        return self.text(box, s, kind="连线说明", fill=WHITE, size=S_SMALL, color=INK, align="c", anchor="m",
                         wrap=False)

    def link(self, a, b, text=None, color=GRAY, lw=1.1, arrow=True, dash=False, shift=(0, 0)):
        """连接两个 Box 的中心，端点裁到矩形边上；text 标在线段中点。"""
        def clip(box, ox, oy):
            dx, dy = ox - box.cx, oy - box.cy
            t = min(box.w / 2 / abs(dx) if dx else 1e9, box.h / 2 / abs(dy) if dy else 1e9)
            return box.cx + dx * t, box.cy + dy * t
        p1, p2 = clip(a, b.cx, b.cy), clip(b, a.cx, a.cy)
        self.line(p1, p2, color, lw, arrow, dash)
        if text:
            self.label((p1[0] + p2[0]) / 2 + shift[0], (p1[1] + p2[1]) / 2 + shift[1], text)

    # ---- 图片、公式、表格
    def figure(self, path, box, align="c", caption=None, kind="图片"):
        """把图片等比缩放到 box 内（不裁剪、不变形）。align 为 c / w / e / n / s。
        caption 给出时在图片正下方放 13pt 图注，图片和图注一起放进 box。返回图片实际占用的 Box。"""
        pw, ph = _png_size(path) if str(path).lower().endswith(".png") else _image_size(path)
        cap_lh = pt2cm(S_SMALL * 1.3)
        cap_h = 0.2 + cap_lh * -(-tw(caption, S_SMALL) // box.w) if caption else 0
        k = min(box.w / pw, (box.h - cap_h) / ph)
        w, h = pw * k, ph * k
        x = {"w": box.l, "e": box.r - w}.get(align, box.cx - w / 2)
        y = {"n": box.t, "s": box.b + h + cap_h}.get(align, box.cy + (h + cap_h) / 2)
        pic = self.shapes.add_picture(str(path), ex(x), ey(y), el(w), el(h))
        self.n += 1
        pic.name = f"{kind} {self.n}"
        if caption:
            self.text(Box(box.l, y - h - 0.2, box.w, cap_h - 0.2), caption, kind="图注", size=S_SMALL,
                      align={"w": "l", "e": "r"}.get(align, "c"), line=S_SMALL * 1.3)
        return Box(x, y, w, h)

    def web_figure(self, path, box, caption, align="c"):
        """放 fetch_image.py 取来的网上图片，图注必填：写图里是什么、图上标号的含义。
        页面上不写出处；来源和许可在旁边同名 .json 里，缺 .json 说明不是按流程取来的图，直接报错。"""
        meta = Path(path).with_suffix(".json")
        if not meta.exists():
            raise FileNotFoundError(f"{meta} 不存在：网上的图要用 fetch_image.py 取，记下来源和许可")
        return self.figure(path, box, align, caption, kind="网图")

    def snippet(self, png, anchor, x, y):
        """按原尺寸放 snippet.py 生成的代码 / 终端图片，锚点同 place。
        返回占用的 Box，并带 line_y(k)：第 k 行（从 0 起，不含标题）文字中心的画布 y 坐标，用于把箭头对准某一行。"""
        meta = json.loads(_png_text(png)["snippet"])
        k = 2.54 / meta["dpi"]
        pw, ph = _png_size(png)
        box = place(anchor, x, y, pw * k, ph * k)
        pic = self.shapes.add_picture(str(png), ex(box.l), ey(box.t), el(box.w), el(box.h))
        self.n += 1
        pic.name = f"截图 {self.n}"
        box.line_y = lambda i: box.t - (meta["head"] + meta["pad"] + meta["lh"] * i + meta["center"]) * k
        return box

    def formula(self, png, x, y, dpi=600):
        """按 dpi 换算原始尺寸放公式图片，(x, y) 为中心。公式用 LaTeX 编译后 pdftoppm -r 600 导出。"""
        pw, ph = _png_size(png)
        w, h = pw / dpi * 2.54, ph / dpi * 2.54
        pic = self.shapes.add_picture(str(png), ex(x - w / 2), ey(y + h / 2), el(w), el(h))
        self.n += 1
        pic.name = f"公式 {self.n}"
        return place("c", x, y, w, h)

    def chart(self, kind, box, categories, series, tones=None, unit=None, number_format="General", labels=False,
              highlight=None, vmax=None):
        """原生图表，数据存在 PPT 内嵌的 Excel 里，可在 PowerPoint 里直接改数。
        kind：bar 竖柱 / barh 横条 / line 折线。series：{系列名: [数值, ...]}，最多 3 个系列，
        颜色按顺序取 tones（默认 primary、neutral、accent：本工作用主色，基线用灰色）。
        unit 写在数值轴标题上；labels=True 时标出数值；highlight 为某个类别名时，单系列柱图里该柱改用强调色。
        vmax 给出时固定数值轴上限（PowerPoint 自动取的上限常留出一截空白）。只有一个系列时不画图例。返回占用的 Box。"""
        if len(series) > 3:
            raise ValueError(f"图表有 {len(series)} 个系列，超过 3 个时拆图或改用表格")
        categories = list(categories)
        data = CategoryChartData(number_format=number_format)
        data.categories = categories
        for name, values in series.items():
            data.add_series(name, values)
        kinds = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "barh": XL_CHART_TYPE.BAR_CLUSTERED,
                 "line": XL_CHART_TYPE.LINE_MARKERS}
        gf = self.shapes.add_chart(kinds[kind], ex(box.l), ey(box.t), el(box.w), el(box.h), data)
        self.n += 1
        gf.name = f"图表 {self.n}"
        ch = gf.chart
        ch.font.size = Pt(S_SMALL)
        ch.font.name = FONT
        ch.font.color.rgb = RGBColor.from_string(TEXT)
        etree.SubElement(ch.font._rPr, qn("a:ea")).set("typeface", FONT)
        ch.has_title = False
        ch.has_legend = len(series) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.TOP
            ch.legend.include_in_layout = False
        plot = ch.plots[0]
        if kind != "line":
            plot.gap_width = 60
        for s, tone in zip(plot.series, tones or ["primary", "neutral", "accent"]):
            color = RGBColor.from_string(TONES[tone].base)
            if kind == "line":
                s.smooth = False
                s.format.line.color.rgb = color
                s.format.line.width = Pt(2.25)
                s.marker.style = XL_MARKER_STYLE.CIRCLE
                s.marker.size = 7
                s.marker.format.fill.solid()
                s.marker.format.fill.fore_color.rgb = color
                s.marker.format.line.color.rgb = color
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = color
                s.format.line.fill.background()
        if highlight is not None:
            point = plot.series[0].points[categories.index(highlight)]
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = RGBColor.from_string(TONES["accent"].base)
        if labels:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = Pt(S_SMALL)
            dl.number_format = number_format
            dl.number_format_is_linked = False
            dl.position = XL_DATA_LABEL_POSITION.ABOVE if kind == "line" else XL_DATA_LABEL_POSITION.OUTSIDE_END
        va, ca = ch.value_axis, ch.category_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = RGBColor.from_string(LINE)
        va.major_gridlines.format.line.width = Pt(0.75)
        va.format.line.fill.background()
        va.tick_labels.number_format = number_format
        va.tick_labels.number_format_is_linked = False
        if kind != "line":  # 柱的长度表示数值，轴不从 0 开始会放大差异
            va.minimum_scale = 0
        if vmax is not None:
            va.maximum_scale = vmax
        if unit:
            va.has_title = True
            fill_text(va.axis_title.text_frame, unit, size=S_SMALL, color=GRAY)
        ca.format.line.color.rgb = RGBColor.from_string(GRAY)
        ca.tick_labels.font.size = Pt(S_SMALL)
        if kind == "barh":
            # 横条图默认第一个类别在最下面；反转类别轴让它排在最上面，和表格的阅读顺序一致。
            # 反转后数值轴会跑到顶部，要在数值轴（不是类别轴）上写 crosses=max 把它放回底部
            ca.reverse_order = True
            va._element.find(qn("c:crosses")).set("val", "max")
        return Box(box.l, box.t, box.w, box.h)

    def table(self, l, t, widths, rows, row_h=1.0, size=S_BODY, header=True, first_col_bold=False):
        """无底色、无网格的原生表格；header 为 True 时首行加粗并在其下画一条横线。
        rows 的每个单元格可以是 str、R、P，或它们的列表（每项一段，用来按意群手动分行）；单元格内文字会按列宽换行。"""
        gf = self.shapes.add_table(len(rows), len(widths), ex(l), ey(t), el(sum(widths)), el(row_h * len(rows)))
        self.n += 1
        gf.name = f"表格 {self.n}"
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        style = tblPr.find(qn("a:tableStyleId"))
        if style is not None:
            tblPr.remove(style)
        tblPr.set("firstRow", "0")
        tblPr.set("bandRow", "0")
        for j, w in enumerate(widths):
            tbl.columns[j].width = el(w)
        for i, row in enumerate(rows):
            tbl.rows[i].height = el(row_h)
            for j, val in enumerate(row):
                cell = tbl.cell(i, j)
                cell.margin_left, cell.margin_right = 0, el(0.2)
                cell.margin_top = cell.margin_bottom = el(0.05)
                is_head = header and i == 0
                fill_text(cell.text_frame, val if isinstance(val, (P, list)) else P(val), size=size,
                          color=INK if is_head or (first_col_bold and j == 0) else TEXT,
                          bold=is_head or (first_col_bold and j == 0), anchor="m")
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                # 不套表格样式时部分软件会画默认网格，四条边逐一写明：表头下画横线，其余行之间画细线
                tcPr = cell._tc.get_or_add_tcPr()
                for side in ("lnL", "lnR", "lnT", "lnB"):
                    if side == "lnB" and is_head:
                        fill, w = f'<a:solidFill><a:srgbClr val="{TEXT}"/></a:solidFill>', 12700
                    elif side == "lnB" and i < len(rows) - 1:
                        fill, w = f'<a:solidFill><a:srgbClr val="{LINE}"/></a:solidFill>', 9525
                    else:
                        fill, w = "<a:noFill/>", 12700
                    tcPr.append(parse_xml(f'<a:{side} {nsdecls("a")} w="{w}">{fill}</a:{side}>'))
        return Box(l, t, sum(widths), row_h * len(rows))


def _png_size(path):
    return struct.unpack(">II", Path(path).read_bytes()[16:24])


def _png_text(path):
    """PNG 里的 tEXt / zTXt 文本块，返回 {关键字: 文本}。"""
    data = Path(path).read_bytes()
    out, pos = {}, 8
    while pos < len(data):
        n = struct.unpack(">I", data[pos:pos + 4])[0]
        kind, body = data[pos + 4:pos + 8], data[pos + 8:pos + 8 + n]
        if kind == b"tEXt":
            key, _, val = body.partition(b"\0")
            out[key.decode("latin-1")] = val.decode("latin-1")
        elif kind == b"zTXt":
            key, _, val = body.partition(b"\0")
            out[key.decode("latin-1")] = zlib.decompress(val[1:]).decode("latin-1")
        pos += 12 + n
    return out


def _image_size(path):
    from pptx.parts.image import Image
    return Image.from_file(str(path)).size


# ============================================================ 版式页

def _set_text(shape, text, size=None):
    """替换文本框文字，保留第一个 run 的格式；size 给出时改字号。
    不用 text_frame.text = ...：它会清掉 run 的字体、颜色等格式。"""
    tf = shape.text_frame
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    runs = tf.paragraphs[0].runs
    runs[0].text = text
    if size:
        runs[0].font.size = Pt(size)
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def _fit_presenter(shape, text):
    """封面「汇报人」框：华科模板里它在组合中，宽约 4.85cm、18pt、字间距 3pt，只放得下 5 个汉字，再长会折行。
    放不下时先去掉字间距再缩字号，最小 14pt；仍放不下就报错，改写组名或第一作者。"""
    grp = shape._element.getparent()
    scale = 1.0
    if grp.tag == qn("p:grpSp"):  # 组合里的形状宽度是组合内坐标，要按组合的缩放换算
        xfrm = grp.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
        scale = int(xfrm.find(qn("a:ext")).get("cx")) / int(xfrm.find(qn("a:chExt")).get("cx"))
    avail = shape.width * scale / CM - 0.5  # 左右各 0.25cm 内边距
    _set_text(shape, text)
    rPr = shape.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
    if tw(text, 18) + pt2cm(3) * len(text) <= avail:
        return
    size = int(avail / tw(text, 1))
    if size < 14:
        raise ValueError(f"汇报人「{text}」太长，封面上一行放不下（14pt 时最多约 8 个汉字），改写组名或第一作者")
    rPr.set("spc", "0")
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(min(size, 18))


def _search(shapes, name):
    for sh in shapes:
        if sh.name == name:
            return sh
        if sh.shape_type == 6:  # 组合
            found = _search(sh.shapes, name)
            if found is not None:
                return found
    return None


def _find(shapes, name):
    found = _search(shapes, name)
    if found is None:
        raise KeyError(f"模板中找不到形状：{name}")
    return found


class Deck:
    """在 templates/<template>/ 的模板上按调用顺序生成幻灯片。"""

    def __init__(self, template):
        self.template = use_template(template)
        self.prs = Presentation(str(self.template.pptx))
        self.cover_t, self.toc_t, self.divider_t, self.content_t = list(self.prs.slides)
        self.sections = []

    def _clone(self, src):
        """复制模板页：重新登记图片等关系再改写 r:id。直接复制幻灯片 XML 文件会让图片关系失效。"""
        new = self.prs.slides.add_slide(src.slide_layout)
        tree = new.shapes._spTree
        for el_ in list(tree)[2:]:
            tree.remove(el_)
        rid_map = {}
        for rid, rel in src.part.rels.items():
            if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
                continue
            target = rel.target_ref if rel.is_external else rel.target_part
            rid_map[rid] = new.part.relate_to(target, rel.reltype, is_external=rel.is_external)
        r_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
        copies = [copy.deepcopy(e) for e in list(src.shapes._spTree)[2:]]
        bg = src._element.cSld.bg
        if bg is not None:
            copies_bg = copy.deepcopy(bg)
            new._element.cSld.insert(0, copies_bg)
            copies.append(copies_bg)
        for c in copies:
            for e in c.iter():
                for k, v in e.attrib.items():
                    if k.startswith(r_ns) and v in rid_map:
                        e.set(k, rid_map[v])
        for c in copies:
            if c.tag != qn("p:bg"):
                tree.append(c)
        return new

    def cover(self, title, presenter, date, title_size=None):
        """封面。标题字号默认 50pt，过长时自动缩小到一行放得下。"""
        s = self._clone(self.cover_t)
        if title_size is None:
            ct = self.template.cover_title
            title_size = min(ct["max_pt"], int(ct["width"] / tw(title, 1)))
        _set_text(_find(s.shapes, "封面标题"), title, title_size)
        _fit_presenter(_find(s.shapes, "汇报人"), presenter)
        _set_text(_find(s.shapes, "日期"), date)
        return s

    def toc(self, sections):
        """目录页：每个章节一条「一、章节名」。过渡页的编号按这里的顺序取。"""
        self.sections = list(sections)
        s = self._clone(self.toc_t)
        proto = _find(s.shapes, "目录条目")._element
        toc = self.template.toc  # 条目从 (x, y) 起竖排，行距 pitch，条目多时压缩到 height 以内
        pitch = min(toc["pitch"], toc["height"] / max(len(sections), 1))
        next_id = max(int(e.get("id")) for e in s._element.iter(qn("p:cNvPr"))) + 1
        for k, sec in enumerate(sections):
            e = copy.deepcopy(proto)
            # 复制出的条目要换新的形状 ID：同一页上 ID 重复时 PowerPoint 拒绝打开整个文件
            e.find(".//" + qn("p:cNvPr")).set("id", str(next_id + k))
            s.shapes._spTree.append(e)
            sh = s.shapes[-1]
            sh.name = f"目录 {k + 1}"
            sh.left, sh.top = el(toc["x"]), el(toc["y"] + k * pitch)
            _set_text(sh, f"{cn_num(k + 1)}、{sec}")
        proto.getparent().remove(proto)
        return s

    def divider(self, name):
        """章节过渡页，编号取该章节在目录中的序号，因此可以只给部分章节加过渡页。"""
        if name not in self.sections:
            raise ValueError(f"过渡页「{name}」不在目录里，先调用 toc()：{self.sections}")
        s = self._clone(self.divider_t)
        _set_text(_find(s.shapes, "章节标题"), f"{cn_num(self.sections.index(name) + 1)}、{name}")
        return s

    def content(self, section, title, notes=""):
        """内容页：右上角章节标签 + 页面标题（写结论句）+ 口播稿备注，返回正文画布。"""
        s = self._clone(self.content_t)
        _set_text(_find(s.shapes, "章节标签"), section)
        _set_text(_find(s.shapes, "页面标题"), title)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return Canvas(s)

    def save(self, path):
        lst = self.prs.slides._sldIdLst
        for tmpl in (self.cover_t, self.toc_t, self.divider_t, self.content_t):
            for sid in list(lst):
                if self.prs.slides.get(int(sid.id)) is tmpl:
                    self.prs.part.drop_rel(sid.rId)
                    lst.remove(sid)
        self.prs.core_properties.category = f"template:{self.template.name}"  # qa_check 据此载入同一套模板
        Path(path).unlink(missing_ok=True)
        self.prs.save(str(path))
